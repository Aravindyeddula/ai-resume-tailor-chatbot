import html
import json
import os
import re
import textwrap
import urllib.request
import ast
from collections import Counter
from difflib import SequenceMatcher
from html.parser import HTMLParser
from pathlib import Path
from typing import Dict, List, Optional, Tuple


TOKEN_RE = re.compile(r"[A-Za-z][A-Za-z0-9+#./-]{1,}")
STOPWORDS = {
    "a",
    "about",
    "across",
    "after",
    "again",
    "against",
    "all",
    "am",
    "an",
    "and",
    "any",
    "are",
    "as",
    "at",
    "be",
    "because",
    "been",
    "before",
    "being",
    "below",
    "between",
    "both",
    "but",
    "by",
    "can",
    "could",
    "did",
    "do",
    "does",
    "doing",
    "down",
    "during",
    "each",
    "few",
    "for",
    "from",
    "further",
    "had",
    "has",
    "have",
    "having",
    "he",
    "her",
    "here",
    "hers",
    "herself",
    "him",
    "himself",
    "his",
    "how",
    "i",
    "if",
    "in",
    "into",
    "is",
    "it",
    "its",
    "itself",
    "just",
    "me",
    "more",
    "most",
    "my",
    "myself",
    "no",
    "nor",
    "not",
    "now",
    "of",
    "off",
    "on",
    "once",
    "only",
    "or",
    "other",
    "our",
    "ours",
    "ourselves",
    "out",
    "over",
    "own",
    "same",
    "she",
    "should",
    "so",
    "some",
    "such",
    "than",
    "that",
    "the",
    "their",
    "theirs",
    "them",
    "themselves",
    "then",
    "there",
    "these",
    "they",
    "this",
    "those",
    "through",
    "to",
    "too",
    "under",
    "until",
    "up",
    "very",
    "was",
    "we",
    "were",
    "what",
    "when",
    "where",
    "which",
    "while",
    "who",
    "whom",
    "why",
    "will",
    "with",
    "you",
    "your",
    "yours",
    "yourself",
    "yourselves",
    "role",
    "position",
    "candidate",
    "responsibilities",
    "requirements",
    "preferred",
    "qualifications",
    "experience",
    "team",
    "teams",
    "ability",
}


class _SimpleHTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: List[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self.parts.append(text)

    def get_text(self) -> str:
        return "\n".join(self.parts)


def _read_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext in {".html", ".htm"}:
        parser = _SimpleHTMLTextExtractor()
        parser.feed(path.read_text(encoding="utf-8", errors="ignore"))
        return html.unescape(parser.get_text())
    if ext == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore
        except Exception:
            return ""
        try:
            reader = PdfReader(str(path))
        except Exception:
            return ""
        parts: List[str] = []
        for page in reader.pages:
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            if text.strip():
                parts.append(text.strip())
        return "\n\n".join(parts)
    if ext == ".docx":
        try:
            from docx import Document  # type: ignore
        except Exception:
            return ""
        try:
            doc = Document(str(path))
        except Exception:
            return ""
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return "\n".join(lines)
    return ""


def _tokenize(text: str) -> List[str]:
    return [t.lower() for t in TOKEN_RE.findall(text)]


def _extract_keywords(text: str, top_n: int = 20) -> List[str]:
    counts = Counter(t for t in _tokenize(text) if t not in STOPWORDS and len(t) >= 3)
    return [w for w, _ in counts.most_common(top_n)]


def _extract_json_object(text: str) -> Dict[str, object]:
    text = text.strip()
    if not text:
        return {}
    try:
        parsed = json.loads(text)
        if isinstance(parsed, dict):
            return parsed
    except Exception:
        pass
    start = text.find("{")
    end = text.rfind("}")
    if start < 0 or end <= start:
        return {}
    try:
        parsed = json.loads(text[start : end + 1])
        if isinstance(parsed, dict):
            return parsed
    except Exception:
        return {}
    return {}


def _key_label(key: object) -> str:
    return str(key).replace("_", " ").strip().title()


def _structured_to_lines(value: object, depth: int = 0) -> List[str]:
    indent = "  " * depth
    lines: List[str] = []
    if isinstance(value, dict):
        for k, v in value.items():
            label = _key_label(k)
            if isinstance(v, (dict, list)):
                lines.append(f"{indent}{label}:")
                lines.extend(_structured_to_lines(v, depth + 1))
            else:
                lines.append(f"{indent}{label}: {str(v)}")
        return lines
    if isinstance(value, list):
        for item in value:
            if isinstance(item, (dict, list)):
                lines.append(f"{indent}-")
                lines.extend(_structured_to_lines(item, depth + 1))
            else:
                lines.append(f"{indent}- {str(item)}")
        return lines
    lines.append(f"{indent}{str(value)}")
    return lines


def _normalize_tailored_text(raw: object) -> str:
    if isinstance(raw, (dict, list)):
        return "\n".join(_structured_to_lines(raw)).strip()
    text = str(raw).strip()
    if not text:
        return ""
    maybe_structured = text.startswith("{") or text.startswith("[")
    if maybe_structured:
        parsed = None
        try:
            parsed = json.loads(text)
        except Exception:
            try:
                parsed = ast.literal_eval(text)
            except Exception:
                parsed = None
        if isinstance(parsed, (dict, list)):
            return "\n".join(_structured_to_lines(parsed)).strip()
    return text


def _build_alignment_section(resume_text: str, jd_text: str) -> Tuple[str, List[str]]:
    jd_keywords = _extract_keywords(jd_text, top_n=18)
    resume_tokens = set(_tokenize(resume_text))

    matched = [k for k in jd_keywords if k in resume_tokens][:6]
    missing = [k for k in jd_keywords if k not in resume_tokens][:8]

    section_lines: List[str] = ["", "Targeted JD Alignment"]
    if matched:
        section_lines.append("- Matched JD keywords: " + ", ".join(matched))
    if missing:
        section_lines.append("- Missing JD keywords to emphasize: " + ", ".join(missing))
    if not matched and not missing:
        section_lines.append("- JD text did not provide enough unique keywords for alignment.")

    summary: List[str] = []
    if matched:
        summary.append(f"Identified {len(matched)} matched JD keywords.")
    if missing:
        summary.append(f"Added {len(missing)} missing JD keywords for targeting.")
    if not summary:
        summary.append("Added JD alignment section.")

    return resume_text.rstrip() + "\n" + "\n".join(section_lines) + "\n", summary


def _extract_evidence_lines(resume_text: str) -> List[str]:
    lines = []
    for raw in resume_text.splitlines():
        s = raw.strip()
        if not s:
            continue
        if len(s) < 30:
            continue
        lines.append(s)
    return lines


def _line_score_for_jd(line: str, jd_keywords: List[str]) -> float:
    tokens = set(_tokenize(line))
    score = 0.0
    overlap = [k for k in jd_keywords if k in tokens]
    score += len(overlap) * 2.0
    if re.search(r"\b\d+(\.\d+)?%?\b", line):
        score += 1.5
    if any(v in line.lower() for v in ["built", "designed", "implemented", "developed", "led", "improved", "automated"]):
        score += 1.0
    return score


def _is_materially_different(original_text: str, updated_text: str) -> bool:
    a = " ".join(original_text.split()).lower()
    b = " ".join(updated_text.split()).lower()
    if a == b:
        return False
    ratio = SequenceMatcher(None, a, b).ratio()
    return ratio < 0.985


def _heuristic_tailor(resume_text: str, jd_text: str) -> Tuple[str, List[str]]:
    jd_keywords = _extract_keywords(jd_text, top_n=18)
    resume_tokens = set(_tokenize(resume_text))
    matched = [k for k in jd_keywords if k in resume_tokens][:8]
    missing = [k for k in jd_keywords if k not in resume_tokens][:8]

    years_match = re.search(r"\b(\d{1,2})\+\s*years\b", resume_text, flags=re.IGNORECASE)
    years = years_match.group(1) + "+" if years_match else "multiple"

    evidence_lines = _extract_evidence_lines(resume_text)
    scored = sorted(
        (( _line_score_for_jd(ln, jd_keywords), ln) for ln in evidence_lines),
        key=lambda x: x[0],
        reverse=True,
    )
    top_evidence = [ln for sc, ln in scored[:6] if sc > 0]
    if not top_evidence:
        top_evidence = evidence_lines[:4]

    summary_keywords = matched[:4] if matched else jd_keywords[:4]
    summary_keyword_text = ", ".join(summary_keywords) if summary_keywords else "data science delivery"

    lines: List[str] = []
    lines.append("TAILORED RESUME SUMMARY")
    lines.append(
        f"Data Scientist with {years} years of experience delivering measurable outcomes in {summary_keyword_text}, "
        "with strong emphasis on production-ready analytics and cross-functional stakeholder impact."
    )
    lines.append("")
    lines.append("JD MATCH SNAPSHOT")
    lines.append("- Matched keywords: " + (", ".join(matched) if matched else "None detected"))
    lines.append("- Keywords to emphasize: " + (", ".join(missing) if missing else "Core JD terms already present"))
    lines.append("")
    lines.append("TARGETED EXPERIENCE HIGHLIGHTS")
    for ev in top_evidence:
        lines.append(f"- {ev}")
    lines.append("")
    lines.append("TARGETED JD ALIGNMENT")
    lines.append("- Focus this version on JD priorities in interviews and profile summary.")

    summary = [
        f"Generated JD-focused summary using {len(summary_keywords)} core terms.",
        f"Selected {len(top_evidence)} experience lines most relevant to the JD.",
    ]
    if missing:
        summary.append(f"Identified {len(missing)} JD keywords to emphasize.")
    return "\n".join(lines).strip() + "\n", summary


def _openai_tailor(resume_text: str, jd_text: str) -> Tuple[str, List[str], str]:
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY not set")
    payload = {
        "model": "gpt-4.1-mini",
        "input": [
            {
                "role": "system",
                "content": (
                    "Rewrite the resume for the provided JD while preserving section order and style. "
                    "Do not invent facts. Return strict JSON with tailored_resume, change_summary (array), notes. "
                    "tailored_resume must be plain resume text (not a dict/object, not markdown code block)."
                ),
            },
            {"role": "user", "content": f"RESUME:\n{resume_text}\n\nJOB_DESCRIPTION:\n{jd_text}"},
        ],
    }
    req = urllib.request.Request(
        "https://api.openai.com/v1/responses",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        body = json.loads(resp.read().decode("utf-8"))

    parsed = _extract_json_object(str(body.get("output_text", "")))
    if not parsed:
        fragments: List[str] = []
        for item in body.get("output", []):
            for content in item.get("content", []):
                if content.get("type") == "output_text" and content.get("text"):
                    fragments.append(str(content["text"]))
        parsed = _extract_json_object("\n".join(fragments))
    if not parsed:
        raise ValueError("Could not parse model output")

    tailored = _normalize_tailored_text(parsed.get("tailored_resume", ""))
    if not tailored:
        raise ValueError("Model returned empty tailored resume")
    raw_summary = parsed.get("change_summary", [])
    if isinstance(raw_summary, list):
        summary = [str(x).strip() for x in raw_summary if str(x).strip()]
    elif isinstance(raw_summary, str):
        summary = [raw_summary.strip()] if raw_summary.strip() else []
    else:
        summary = []
    notes = str(parsed.get("notes", "")).strip()
    return tailored, summary, notes


def _render_html_document(text: str) -> str:
    return (
        "<!doctype html><html><head><meta charset='utf-8'><title>Tailored Resume</title></head>"
        "<body><pre style='white-space:pre-wrap;'>"
        + html.escape(text)
        + "</pre></body></html>\n"
    )


def _chunk_lines_for_blocks(all_lines: List[str], capacities: List[int]) -> List[str]:
    chunks: List[str] = []
    cursor = 0
    for cap in capacities:
        size = max(1, int(cap))
        selected = all_lines[cursor : cursor + size]
        cursor += size
        chunks.append("\n".join(selected).strip())
    if cursor < len(all_lines) and chunks:
        extra = "\n".join(all_lines[cursor:]).strip()
        if extra:
            chunks[-1] = (chunks[-1] + "\n" + extra).strip()
    return chunks


def _prepare_layout_lines(text: str, width: int = 90) -> List[str]:
    out: List[str] = []
    for raw in text.splitlines():
        s = raw.strip()
        if not s:
            continue
        wrapped = textwrap.wrap(s, width=width, break_long_words=False, break_on_hyphens=False)
        if wrapped:
            out.extend(wrapped)
        else:
            out.append(s)
    return out


def _write_docx_output(text: str, path: Path, template_path: Optional[Path] = None) -> Optional[str]:
    try:
        from docx import Document  # type: ignore
    except Exception:
        return "Missing dependency: python-docx."
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if template_path and template_path.exists():
        try:
            doc = Document(str(template_path))
            targets = [p for p in doc.paragraphs if p.text.strip()]
            if not targets:
                targets = list(doc.paragraphs)
            for i, p in enumerate(targets):
                p.text = lines[i] if i < len(lines) else ""
            if len(lines) > len(targets):
                style = targets[-1].style if targets else None
                for ln in lines[len(targets) :]:
                    np = doc.add_paragraph(ln)
                    if style is not None:
                        np.style = style
        except Exception:
            doc = Document()
            for line in lines or [""]:
                doc.add_paragraph(line)
    else:
        doc = Document()
        for line in lines or [""]:
            doc.add_paragraph(line)
    doc.save(str(path))
    return None


def _normalize_pdf_color(color_val: object) -> Tuple[float, float, float]:
    if isinstance(color_val, int):
        rgb = int(color_val) & 0xFFFFFF
        r = ((rgb >> 16) & 255) / 255.0
        g = ((rgb >> 8) & 255) / 255.0
        b = (rgb & 255) / 255.0
        return (r, g, b)
    return (0.0, 0.0, 0.0)


def _safe_pdf_fontname(raw_font: str) -> str:
    f = (raw_font or "").strip()
    if "+" in f:
        f = f.split("+", 1)[1]
    return f or "Helvetica"


def _resolve_windows_font(font_name: str) -> Tuple[str, Optional[str]]:
    low = (font_name or "").lower()
    windows_font_dir = Path(r"C:\Windows\Fonts")

    def _pick(path_name: str, alias: str) -> Tuple[str, Optional[str]]:
        p = windows_font_dir / path_name
        if p.exists():
            return alias, str(p)
        return "helv", None

    if "calibri" in low:
        if "bold" in low and ("italic" in low or "oblique" in low):
            return _pick("calibriz.ttf", "FCalBI")
        if "bold" in low:
            return _pick("calibrib.ttf", "FCalB")
        if "italic" in low or "oblique" in low:
            return _pick("calibrii.ttf", "FCalI")
        return _pick("calibri.ttf", "FCal")
    if "arial" in low:
        if "bold" in low and ("italic" in low or "oblique" in low):
            return _pick("arialbi.ttf", "FArialBI")
        if "bold" in low:
            return _pick("arialbd.ttf", "FArialB")
        if "italic" in low or "oblique" in low:
            return _pick("ariali.ttf", "FArialI")
        return _pick("arial.ttf", "FArial")
    if "times" in low:
        if "bold" in low and ("italic" in low or "oblique" in low):
            return _pick("timesbi.ttf", "FTimesBI")
        if "bold" in low:
            return _pick("timesbd.ttf", "FTimesB")
        if "italic" in low or "oblique" in low:
            return _pick("timesi.ttf", "FTimesI")
        return _pick("times.ttf", "FTimes")
    return "helv", None


def _extract_pdf_style_blocks(doc: object) -> List[Dict[str, object]]:
    blocks: List[Dict[str, object]] = []
    try:
        import fitz  # type: ignore
    except Exception:
        return blocks

    for page_index in range(doc.page_count):
        page = doc[page_index]
        text_dict = page.get_text("dict")
        for block in text_dict.get("blocks", []):
            if block.get("type") != 0:
                continue
            lines = block.get("lines", [])
            if not lines:
                continue
            merged_lines: List[str] = []
            first_span = None
            for ln in lines:
                span_texts: List[str] = []
                for span in ln.get("spans", []):
                    txt = str(span.get("text", ""))
                    if txt:
                        span_texts.append(txt)
                    if first_span is None and txt.strip():
                        first_span = span
                if span_texts:
                    merged_lines.append("".join(span_texts).strip())
            if not merged_lines:
                continue
            if first_span is None:
                first_span = {"font": "helv", "size": 11, "color": 0}
            blocks.append(
                {
                    "page": page_index,
                    "rect": fitz.Rect(block["bbox"]),
                    "line_capacity": max(1, len(merged_lines)),
                    "font": _safe_pdf_fontname(str(first_span.get("font", "helv"))),
                    "size": float(first_span.get("size", 11)),
                    "color": _normalize_pdf_color(first_span.get("color", 0)),
                }
            )
    blocks.sort(key=lambda b: (int(b["page"]), float(b["rect"].y0), float(b["rect"].x0)))
    return blocks


def _extract_pdf_typography(template_path: Path) -> Tuple[float, float, str, float, Tuple[float, float, float]]:
    try:
        import fitz  # type: ignore
    except Exception:
        return 612.0, 792.0, "Helvetica", 10.0, (0.0, 0.0, 0.0)

    doc = fitz.open(str(template_path))
    try:
        first = doc[0]
        width = float(first.rect.width)
        height = float(first.rect.height)
        font_weights: Dict[str, int] = {}
        size_weights: Dict[float, int] = {}
        color_weights: Dict[Tuple[float, float, float], int] = {}

        for page in doc:
            td = page.get_text("dict")
            for block in td.get("blocks", []):
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        txt = str(span.get("text", ""))
                        if not txt.strip():
                            continue
                        w = len(txt)
                        f = _safe_pdf_fontname(str(span.get("font", "Helvetica")))
                        s = round(float(span.get("size", 10.0)), 2)
                        c = _normalize_pdf_color(span.get("color", 0))
                        font_weights[f] = font_weights.get(f, 0) + w
                        size_weights[s] = size_weights.get(s, 0) + w
                        color_weights[c] = color_weights.get(c, 0) + w

        dominant_font = max(font_weights.items(), key=lambda x: x[1])[0] if font_weights else "Helvetica"
        dominant_size = max(size_weights.items(), key=lambda x: x[1])[0] if size_weights else 10.0
        dominant_color = max(color_weights.items(), key=lambda x: x[1])[0] if color_weights else (0.0, 0.0, 0.0)
        return width, height, dominant_font, float(dominant_size), dominant_color
    finally:
        doc.close()


def _write_pdf_output(text: str, path: Path, template_path: Optional[Path] = None) -> Optional[str]:
    lines = _prepare_layout_lines(text, width=95) or [""]

    # Preferred strategy: generate tailored-only PDF using template page-size and style cues.
    if template_path and template_path.exists():
        try:
            import fitz  # type: ignore

            width, height, dominant_font, dominant_size, dominant_color = _extract_pdf_typography(template_path)
            base_font, base_font_file = _resolve_windows_font(dominant_font)
            base_size = max(9.0, float(dominant_size))
            base_color = dominant_color

            out_doc = fitz.open()
            margin_x = 46.0
            margin_top = 46.0
            margin_bottom = 46.0
            body_line_h = max(11.0, base_size + 2.0)
            y = margin_top

            page = out_doc.new_page(width=width, height=height)
            if base_font_file:
                try:
                    page.insert_font(fontname=base_font, fontfile=base_font_file)
                except Exception:
                    base_font = "helv"
                    base_font_file = None
            page.insert_text(
                (margin_x, y),
                "Tailored Resume",
                fontname=base_font,
                fontsize=max(base_size + 3.0, 13.0),
                color=base_color,
            )
            y += body_line_h * 1.6

            for ln in lines:
                wrapped = textwrap.wrap(ln, width=92, break_long_words=False, break_on_hyphens=False) or [ln]
                for part in wrapped:
                    if y > height - margin_bottom:
                        page = out_doc.new_page(width=width, height=height)
                        if base_font_file:
                            try:
                                page.insert_font(fontname=base_font, fontfile=base_font_file)
                            except Exception:
                                base_font = "helv"
                                base_font_file = None
                        y = margin_top
                    try:
                        page.insert_text(
                            (margin_x, y),
                            part,
                            fontname=base_font,
                            fontsize=base_size,
                            color=base_color,
                        )
                    except Exception:
                        page.insert_text(
                            (margin_x, y),
                            part,
                            fontname="helv",
                            fontsize=base_size,
                            color=base_color,
                        )
                    y += body_line_h

            out_doc.save(str(path))
            out_doc.close()
            return "Generated tailored-only PDF using template page size and style cues."
        except Exception as e:
            template_warning = f"Could not render tailored-only template PDF: {e}"
    else:
        template_warning = None

    # If template-based append is unavailable, render a standalone readable PDF.
    try:
        from reportlab.lib.pagesizes import LETTER  # type: ignore
        from reportlab.pdfgen import canvas  # type: ignore
    except Exception:
        return "Missing dependencies: pymupdf/reportlab."

    width, height = LETTER
    margin_x = 50
    margin_y = 50
    line_height = 14
    y = height - margin_y
    c = canvas.Canvas(str(path), pagesize=LETTER)
    c.setTitle("Tailored Resume")
    c.setFont("Times-Roman", 11)
    for raw_line in lines:
        wrapped = textwrap.wrap(raw_line, width=95, break_long_words=False, break_on_hyphens=False) or [raw_line]
        for line in wrapped:
            c.drawString(margin_x, y, line)
            y -= line_height
            if y <= margin_y:
                c.showPage()
                c.setFont("Times-Roman", 11)
                y = height - margin_y
    c.save()
    return template_warning or "Used generic PDF writer (template not available)."


def _write_output(
    text: str,
    output_dir: Path,
    stem: str,
    ext: str,
    template_path: Optional[Path] = None,
) -> Tuple[Path, str, Optional[str]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    if ext in {".txt", ".md"}:
        out = output_dir / f"{stem}_tailored{ext}"
        out.write_text(text, encoding="utf-8")
        return out, ext, None
    if ext in {".html", ".htm"}:
        out = output_dir / f"{stem}_tailored{ext}"
        out.write_text(_render_html_document(text), encoding="utf-8")
        return out, ext, None
    if ext == ".docx":
        out = output_dir / f"{stem}_tailored.docx"
        warning = _write_docx_output(text, out, template_path=template_path)
        if out.exists():
            return out, ".docx", warning
        if warning is None:
            return out, ".docx", None
        txt_out = output_dir / f"{stem}_tailored.txt"
        txt_out.write_text(text, encoding="utf-8")
        return txt_out, ".txt", warning
    if ext == ".pdf":
        out = output_dir / f"{stem}_tailored.pdf"
        warning = _write_pdf_output(text, out, template_path=template_path)
        if out.exists():
            return out, ".pdf", warning
        if warning is None:
            return out, ".pdf", None
        txt_out = output_dir / f"{stem}_tailored.txt"
        txt_out.write_text(text, encoding="utf-8")
        return txt_out, ".txt", warning
    out = output_dir / f"{stem}_tailored.txt"
    out.write_text(text, encoding="utf-8")
    return out, ".txt", "Unsupported input extension; exported as .txt."


def _tailor_from_text(
    resume_text: str,
    jd_text: str,
    output_dir: Path,
    stem: str,
    preferred_ext: str,
    source_path: Optional[Path] = None,
) -> Dict[str, object]:
    cleaned_resume = resume_text.strip()
    cleaned_jd = jd_text.strip()
    if not cleaned_resume:
        return {"ok": False, "error": "Resume text is empty."}
    if not cleaned_jd:
        return {"ok": False, "error": "JD text is empty."}

    mode = "heuristic"
    notes = ""
    try:
        tailored, summary, notes = _openai_tailor(cleaned_resume, cleaned_jd)
        mode = "openai"
    except Exception as e:
        tailored, summary = _heuristic_tailor(cleaned_resume, cleaned_jd)
        notes = f"Fallback used: {e}"

    if not _is_materially_different(cleaned_resume, tailored) and "Targeted JD Alignment" not in tailored:
        tailored, forced_summary = _build_alignment_section(cleaned_resume, cleaned_jd)
        summary = list(summary) + forced_summary
        notes = f"{notes} | Added deterministic JD alignment because rewrite was too similar.".strip(" |")

    output_path, output_ext, writer_warning = _write_output(
        tailored,
        output_dir,
        stem,
        preferred_ext.lower(),
        template_path=source_path,
    )
    if writer_warning:
        notes = f"{notes} | {writer_warning}".strip(" |")
    return {
        "ok": True,
        "mode": mode,
        "output_path": output_path.as_posix(),
        "output_ext": output_ext,
        "format_preserved": output_ext == preferred_ext.lower(),
        "change_summary": summary,
        "notes": notes,
    }


def tailor_resume(resume_path: str, jd_path: str, output_dir: str = "artifacts") -> Dict[str, object]:
    resume_file = Path(resume_path)
    jd_file = Path(jd_path)
    if not resume_file.exists():
        return {"ok": False, "error": f"Resume file not found: {resume_file.as_posix()}"}
    if not jd_file.exists():
        return {"ok": False, "error": f"JD file not found: {jd_file.as_posix()}"}

    resume_text = _read_file(resume_file).strip()
    jd_text = _read_file(jd_file).strip()
    if not resume_text:
        return {"ok": False, "error": "Could not extract resume text."}
    if not jd_text:
        return {"ok": False, "error": "Could not extract JD text."}

    return _tailor_from_text(
        resume_text=resume_text,
        jd_text=jd_text,
        output_dir=Path(output_dir),
        stem=resume_file.stem,
        preferred_ext=resume_file.suffix.lower(),
        source_path=resume_file,
    )


def tailor_resume_text(
    resume_text: str,
    jd_text: str,
    output_dir: str = "artifacts",
    output_stem: str = "pasted_resume",
    preferred_ext: str = ".txt",
) -> Dict[str, object]:
    return _tailor_from_text(
        resume_text=resume_text,
        jd_text=jd_text,
        output_dir=Path(output_dir),
        stem=output_stem,
        preferred_ext=preferred_ext,
        source_path=None,
    )


def tailor_resume_with_jd_text(
    resume_path: str,
    jd_text: str,
    output_dir: str = "artifacts",
) -> Dict[str, object]:
    resume_file = Path(resume_path)
    if not resume_file.exists():
        return {"ok": False, "error": f"Resume file not found: {resume_file.as_posix()}"}

    resume_text = _read_file(resume_file).strip()
    if not resume_text:
        return {"ok": False, "error": "Could not extract resume text from uploaded file."}

    return _tailor_from_text(
        resume_text=resume_text,
        jd_text=jd_text,
        output_dir=Path(output_dir),
        stem=resume_file.stem,
        preferred_ext=resume_file.suffix.lower(),
        source_path=resume_file,
    )
